from docx import Document
from copy import deepcopy
import requests
import os
import bs4
from openpyxl import load_workbook
import pandas as pd
from ..helpers.db_funcs import get_ep_id, get_ep_id_by_number, get_season_id_by_number_type
from ..helpers.extract_helpers import search_for_new_seasons
import glob
import re
from pydrive.auth import GoogleAuth
from pydrive.drive import GoogleDrive
import numpy as np

season_type_map = {
    's': 'Survivor',
    'au': 'Australian Survivor',
    'nz': 'Survivor New Zealand',
    'ssa': 'Survivor South Africa'
}


def dfize_doc(doc):
    df = pd.DataFrame()
    day = None
    order = 1
    extracted_names = ['contestant', 'n_from_player',
                       'total_confessionals_in_episode', 'content', 'day', 'n_in_episode']
    for p in doc.paragraphs:
        day_match = re.match('Day (\d+)', p.text)
        conf_match = re.match('(\w+ ?\w+) \((\d+)\/(\d+)\)\: (.*)', p.text)

        if day_match:
            day = day_match.group(1)

        if day:
            if conf_match:
                extracted = conf_match.groups(0)
                srs = pd.Series(list(extracted) +
                                [day, order], index=extracted_names)
                df = df.append(srs, ignore_index=True)
                order += 1

    return df


def collect_confessionals(raw_files, eng):
    full_df = pd.DataFrame()
    regex_file = '(\d+)x(\d+)_ (.*) (\(Pt\. \d\))?\.docx'
    for f in raw_files:
        basename = os.path.basename(f)
        m = re.match(regex_file, basename)
        if m:
            season, episode, episode_name, _ = m.groups(0)
            doc = Document(f)
            df = dfize_doc(doc)
            season_id = get_season_id_by_number_type(
                eng, int(season), 'Survivor')
            episode_id = get_ep_id(eng, episode_name, season_id)

            df['season_id'] = season_id
            df['episode_id'] = episode_id
            full_df = pd.concat([full_df, df], ignore_index=True)

    return full_df


def extract_confessionals(eng,
                          data_path=None,
                          asof=None):
    if data_path is None:
        data_path = os.path.join(os.path.dirname(__file__),
                                 '../../../data/raw/confessionals')
    sync_confessionals(data_path)
    raw_files = glob.glob(os.path.join(data_path, '*', '*'))
    new_seasons = search_for_new_seasons(eng, asof=asof)

    new_seasons_df = pd.read_sql(con=eng,
                                 sql='SELECT * FROM survivor.season').set_index('name').loc[new_seasons]
    search_string = new_seasons_df['season_number'].astype(
        int).astype(str).str.zfill(2) + 'x'
    print(search_string)

    raw_files = [f for f in raw_files if any(
        re.search(y, f) for y in search_string)]
    print(raw_files)

    full_df = collect_confessionals(raw_files, eng)
    return full_df


# Syncing with Google Drive (for Confessional Data)
MIMETYPES = {
    'application/vnd.google-apps.document': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    'application/vnd.google-apps.spreadsheet': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
}


def download_special_file(file, output_filename):
    if file['mimeType'] in MIMETYPES:
        download_mimetype = MIMETYPES[file['mimeType']]
        file.GetContentFile(output_filename, mimetype=download_mimetype)

    else:
        file.GetContentFile(output_filename)

    # Work around, closing file
    for c in file.http.connections.values():
        c.close()


def sync_confessionals(data_dir='test'):
    if not os.path.exists(data_dir):
        os.mkdir(data_dir)

    GoogleAuth.DEFAULT_SETTINGS['client_config_file'] = os.path.abspath(os.path.join(os.path.dirname(
        __file__), '../../../../../client_config.json'))
    gauth = GoogleAuth()
    gauth.LoadCredentialsFile('google_drive_creds.txt')

    drive = GoogleDrive(gauth)

    gd_files = pull_confessionals_files(drive)

    for subfolder, file_list in gd_files.items():
        if not os.path.exists(os.path.join(data_dir, subfolder)):
            os.mkdir(os.path.join(data_dir, subfolder))
        data_files = os.listdir(os.path.join(data_dir, subfolder))

        for f in file_list:

            title = f['title']
            full_path = os.path.join(data_dir, subfolder, title) + '.docx'
            if title not in data_files:
                download_special_file(f, full_path)
            else:
                m_date = os.path.getmtime(full_path)
                print(f['modifiedDate'])
                if pd.to_datetime(f['modifiedDate']) > pd.to_datetime(m_date, utc=True, unit='s'):
                    # Overwrite it
                    download_special_file(f, full_path)
                    print('Overwrote the file with a newer version')
                else:
                    print('Newest file already in local directory')


def pull_confessionals_files(drive):

    query = """title contains 'Archive'
               and 'Ismael' in owners
               and title contains 'Survivor Confessional'
               and mimeType = 'application/vnd.google-apps.folder'"""

    folder_id = drive.ListFile(dict(q=query)).GetList()[0]['id']

    subfolder_query = """'{folder_id}' in parents
                           and mimeType = 'application/vnd.google-apps.folder'""".format(folder_id=folder_id)

    subfolders = drive.ListFile(dict(q=subfolder_query)).GetList()

    subfolders_to_files = {}
    for subfolder in subfolders:
        files_query = "'{folder_id}' in parents".format(
            folder_id=subfolder['id'])
        files = drive.ListFile(dict(q=files_query)).GetList()

        subfolders_to_files[subfolder['title']] = files

    return subfolders_to_files
